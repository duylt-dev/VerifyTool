"""Read real workbooks in MO's layout (static-check-spec §3)."""

from pathlib import Path

import openpyxl
import pytest
from conftest import HEADERS, UNIT_ID, make_spec_xlsx

from apkverify.errors import SpecError
from apkverify.input.spec_reader import read_spec


@pytest.mark.parametrize("nfd_header", [False, True])
def test_layout_and_group_placement_row_18(tmp_path: Path, nfd_header: bool) -> None:
    rows = [
        {"Định Dạng": "Native", "Config Name": " before ", "Ad Space": " Before "},
        {"STT": 1, "Task": " Language ", "Ghi chú AD": " group note "},
        {
            "Định Dạng": "Native",
            "Config Name": "one",
            "Ad Space": "one_native",
            "ID": UNIT_ID,
            "Task": " placement ",
            "Trạng thái": "done",
            "Network Ad Name": "Am_one_native",
            "Ghi chú AD": "own",
            "Mã task": "MO-1",
            "Link ClickUp": "https://example.com/task",
        },
    ]
    rows.extend({} for _ in range(10))
    rows.append(
        {
            "STT": 2,
            "Task": "QC native full sau inter",
            "Ghi chú AD": "new group",
            "Định Dạng": "Native Full Screen",
            "Config Name": "afterinterstitial",
            "Ad Space": "after_native",
            "ID": UNIT_ID,
        }
    )
    spec = read_spec(make_spec_xlsx(tmp_path / "spec.xlsx", rows, nfd_header=nfd_header))
    assert (spec.file_name, spec.sheet, spec.title) == ("spec.xlsx", "Ads", "Quảng Cáo / Ads_Test")
    before, placement, both = spec.placements
    assert (before.row, before.group, before.group_note, before.config_name, before.space_name) == (
        5,
        "",
        "",
        "before",
        "Before",
    )
    assert (placement.row, placement.group, placement.group_note, placement.task) == (
        7,
        "Language",
        "group note",
        "placement",
    )
    assert (placement.status, placement.network_name, placement.note, placement.task_code, placement.link) == (
        "done",
        "Am_one_native",
        "own",
        "MO-1",
        "https://example.com/task",
    )
    assert both.row == 18 and both.group == both.task == "QC native full sau inter"
    assert both.group_note == both.note == "new group"


def test_integral_float_and_first_sheet(tmp_path: Path) -> None:
    path = make_spec_xlsx(tmp_path / "spec.xlsx", [{"Định Dạng": "Native", "Config Name": 12.0, "ID": 7.5}])
    workbook = openpyxl.load_workbook(path)
    workbook.create_sheet("Ignored").append(["other"])
    workbook.save(path)
    workbook.close()
    placement = read_spec(path).placements[0]
    assert placement.config_name == "12" and placement.ad_id == "7.5"
    assert placement.group == placement.note == ""


@pytest.mark.parametrize("header_row", [1, 20, 21])
def test_header_search_boundary_and_reordered_columns(tmp_path: Path, header_row: int) -> None:
    path = tmp_path / "spec.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    headers = list(reversed(HEADERS))
    for index, header in enumerate(headers, 1):
        sheet.cell(header_row, index, " " + header.swapcase() + " ")
    sheet.cell(header_row + 1, headers.index("Định Dạng") + 1, "Native")
    sheet.cell(header_row + 1, headers.index("Config Name") + 1, "home")
    workbook.save(path)
    workbook.close()
    if header_row == 21:
        with pytest.raises(SpecError) as error:
            read_spec(path)
        assert str(error.value) == (
            "Không tìm thấy dòng tiêu đề trong 20 dòng đầu của spec.xlsx "
            "(cần các cột: Định Dạng, Config Name, Ad Space, ID)."
        )
    else:
        spec = read_spec(path)
        assert spec.title == ""
        assert spec.placements[0].row == header_row + 1
        assert spec.placements[0].config_name == "home"


def test_missing_header(tmp_path: Path) -> None:
    path = tmp_path / "empty.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["Title"])
    workbook.save(path)
    workbook.close()
    with pytest.raises(SpecError) as error:
        read_spec(path)
    assert str(error.value) == (
        "Không tìm thấy dòng tiêu đề trong 20 dòng đầu của empty.xlsx "
        "(cần các cột: Định Dạng, Config Name, Ad Space, ID)."
    )


def test_unreadable_workbook(tmp_path: Path) -> None:
    path = tmp_path / "broken.xlsx"
    path.write_text("not a workbook", encoding="utf-8")
    with pytest.raises(SpecError) as error:
        read_spec(path)
    assert str(error.value) == "Không đọc được file spec broken.xlsx: File is not a zip file."
